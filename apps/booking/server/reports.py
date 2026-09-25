"""分析與報表：期間統計與 Excel 匯出。

期間以日期（含頭尾）計算：
- 營收：期間內「確認收款」的訂單（以確認時間計）
- 上課：課程日期落在期間內、未停課的課程
- 課卡消耗：上述課程中實際扣掉的堂數／點數
"""
import io
from datetime import date, datetime, timedelta

BOOKED = "('booked','attended','absent')"


def _rows(conn, sql, args=()):
    return [dict(r) for r in conn.execute(sql, args).fetchall()]


def _days(start: str, end: str) -> list[str]:
    d, last, out = date.fromisoformat(start), date.fromisoformat(end), []
    while d <= last:
        out.append(d.isoformat())
        d += timedelta(days=1)
    return out


def _minutes(c: dict) -> int:
    s = datetime.fromisoformat(f"2000-01-01T{c['start_time']}")
    e = datetime.fromisoformat(f"2000-01-01T{c['end_time']}")
    return max(int((e - s).total_seconds() // 60), 0)


def _pct(a, b):
    return round(a / b * 100) if b else None


def summary(conn, start: str, end: str, today: str) -> dict:
    # ---------- 營收
    orders = _rows(conn,
                   "SELECT o.*, p.name plan_name, p.type plan_type, u.name user_name, u.phone FROM orders o"
                   " JOIN plans p ON p.id=o.plan_id JOIN users u ON u.id=o.user_id"
                   " WHERE o.status='paid' AND substr(o.updated_at,1,10) BETWEEN ? AND ? ORDER BY o.updated_at",
                   (start, end))
    by_plan: dict[str, dict] = {}
    daily_rev = {d: 0 for d in _days(start, end)}
    for o in orders:
        p = by_plan.setdefault(o["plan_name"], {"name": o["plan_name"], "type": o["plan_type"], "count": 0, "amount": 0})
        p["count"] += 1
        p["amount"] += o["amount"]
        daily_rev[o["updated_at"][:10]] = daily_rev.get(o["updated_at"][:10], 0) + o["amount"]

    # ---------- 上課
    courses = _rows(conn, "SELECT c.*, t.name teacher_name FROM courses c LEFT JOIN teachers t ON t.id=c.teacher_id"
                          " WHERE c.date BETWEEN ? AND ? ORDER BY c.date, c.start_time", (start, end))
    open_courses = [c for c in courses if c["status"] == "open"]
    ids = [c["id"] for c in open_courses] or [0]
    marks = ",".join("?" * len(ids))
    res = _rows(conn, f"SELECT r.*, u.name user_name, u.phone FROM reservations r JOIN users u ON u.id=r.user_id"
                      f" WHERE r.course_id IN ({marks})", ids)
    by_course: dict[int, list] = {}
    for r in res:
        by_course.setdefault(r["course_id"], []).append(r)
    daily_book = {d: 0 for d in _days(start, end)}
    capacity = booked = attended = absent = cancelled = 0
    course_rank: dict[str, dict] = {}
    teacher_rank: dict[str, dict] = {}
    for c in open_courses:
        rs = by_course.get(c["id"], [])
        b = sum(r["status"] in ("booked", "attended", "absent") for r in rs)
        a = sum(r["status"] == "attended" for r in rs)
        ab = sum(r["status"] == "absent" for r in rs)
        capacity += c["capacity"]
        booked += b
        attended += a
        absent += ab
        cancelled += sum(r["status"] == "cancelled" for r in rs)
        daily_book[c["date"]] = daily_book.get(c["date"], 0) + b
        cr = course_rank.setdefault(c["name"], {"name": c["name"], "sessions": 0, "capacity": 0, "booked": 0, "attended": 0})
        cr["sessions"] += 1
        cr["capacity"] += c["capacity"]
        cr["booked"] += b
        cr["attended"] += a
        tname = c["teacher_name"] or "未指定"
        tr = teacher_rank.setdefault(tname, {"name": tname, "sessions": 0, "minutes": 0, "booked": 0, "attended": 0})
        tr["sessions"] += 1
        tr["minutes"] += _minutes(c)
        tr["booked"] += b
        tr["attended"] += a
    for cr in course_rank.values():
        cr["fill_rate"] = _pct(cr["booked"], cr["capacity"])
    for tr in teacher_rank.values():
        tr["hours"] = round(tr["minutes"] / 60, 1)

    # ---------- 課卡
    consumed = {"sessions": 0, "points": 0, "unlimited": 0}
    for r in res:
        if r["status"] in ("booked", "attended", "absent") and r["card_id"]:
            card = conn.execute("SELECT type FROM cards WHERE id=?", (r["card_id"],)).fetchone()
            if card:
                consumed[card["type"]] += r["charged"] if card["type"] != "unlimited" else 1
    outstanding = {row["type"]: {"cards": row["n"], "remaining": row["rem"]} for row in conn.execute(
        "SELECT type, COUNT(*) n, SUM(remaining) rem FROM cards WHERE expires_on>=?"
        " AND (type='unlimited' OR remaining>0) GROUP BY type", (today,))}
    manual = conn.execute("SELECT COUNT(*) FROM cards WHERE source='manual' AND substr(created_at,1,10) BETWEEN ? AND ?",
                          (start, end)).fetchone()[0]

    # ---------- 會員
    new_members = conn.execute("SELECT COUNT(*) FROM users WHERE role='student' AND substr(created_at,1,10) BETWEEN ? AND ?",
                               (start, end)).fetchone()[0]
    member_rank: dict[int, dict] = {}
    for r in res:
        if r["status"] in ("booked", "attended", "absent"):
            m = member_rank.setdefault(r["user_id"], {"name": r["user_name"], "phone": r["phone"], "booked": 0,
                                                      "attended": 0, "absent": 0})
            m["booked"] += 1
            m["attended"] += r["status"] == "attended"
            m["absent"] += r["status"] == "absent"

    return {
        "from": start, "to": end,
        "revenue": {
            "total": sum(o["amount"] for o in orders), "orders": len(orders),
            "by_plan": sorted(by_plan.values(), key=lambda x: -x["amount"]),
            "daily": [{"date": d, "value": v} for d, v in daily_rev.items()],
            "pending": conn.execute("SELECT COUNT(*) FROM orders WHERE status='pending'").fetchone()[0],
        },
        "classes": {
            "sessions": len(open_courses), "cancelled_sessions": len(courses) - len(open_courses),
            "capacity": capacity, "booked": booked, "fill_rate": _pct(booked, capacity),
            "attended": attended, "absent": absent, "attendance_rate": _pct(attended, attended + absent),
            "cancellations": cancelled,
            "daily": [{"date": d, "value": v} for d, v in daily_book.items()],
        },
        "courses": sorted(course_rank.values(), key=lambda x: (-x["booked"], x["name"])),
        "teachers": sorted(teacher_rank.values(), key=lambda x: (-x["sessions"], x["name"])),
        "cards": {"consumed": consumed, "outstanding": outstanding, "manual_grants": manual},
        "members": {
            "new": new_members, "active": len(member_rank),
            "top": sorted(member_rank.values(), key=lambda x: (-x["booked"], x["name"]))[:10],
        },
        "_orders": orders, "_courses": open_courses, "_res": res,
    }


STATUS = {"booked": "已預約", "attended": "出席", "absent": "缺席", "waitlist": "候補", "cancelled": "已取消"}
TYPES = {"sessions": "堂數課卡", "points": "點數課卡", "unlimited": "無限課卡"}


def to_xlsx(data: dict, venue: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1A6FE0")

    def sheet(title, header, body, widths=None, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center")
        for row in body:
            ws.append(row)
        for i, w in enumerate(widths or [14] * len(header)):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = "A2"
        return ws

    rev, cls, cards, mem = data["revenue"], data["classes"], data["cards"], data["members"]
    sheet("摘要", ["項目", "數值"], [
        ["場館", venue], ["期間", f"{data['from']} ~ {data['to']}"],
        ["營收（NT$）", rev["total"]], ["已收款訂單", rev["orders"]], ["待確認訂單（目前）", rev["pending"]],
        ["開課堂數", cls["sessions"]], ["停課堂數", cls["cancelled_sessions"]], ["總名額", cls["capacity"]],
        ["預約人次", cls["booked"]], ["滿班率（%）", cls["fill_rate"]], ["出席", cls["attended"]],
        ["缺席", cls["absent"]], ["出席率（%）", cls["attendance_rate"]], ["取消預約", cls["cancellations"]],
        ["新會員", mem["new"]], ["上課會員", mem["active"]],
        ["消耗堂數", cards["consumed"]["sessions"]], ["消耗點數", cards["consumed"]["points"]],
        ["無限卡上課次數", cards["consumed"]["unlimited"]], ["場主手動開通課卡", cards["manual_grants"]],
        *[[f"未使用{TYPES[k]}（{'張' if k == 'unlimited' else '剩餘'}）",
           v["cards"] if k == "unlimited" else v["remaining"]] for k, v in cards["outstanding"].items()],
    ], [26, 26], first=True)
    sheet("營收明細", ["收款日期", "會員", "手機", "方案", "類型", "金額"],
          [[o["updated_at"][:16].replace("T", " "), o["user_name"], o["phone"], o["plan_name"],
            TYPES.get(o["plan_type"], ""), o["amount"]] for o in data["_orders"]], [18, 12, 14, 16, 12, 10])
    courses = {c["id"]: c for c in data["_courses"]}
    res = sorted(data["_res"], key=lambda r: (courses[r["course_id"]]["date"], courses[r["course_id"]]["start_time"]))
    sheet("上課明細", ["日期", "時間", "課程", "老師", "學員", "手機", "狀態", "扣除"],
          [[courses[r["course_id"]]["date"], f"{courses[r['course_id']]['start_time']}~{courses[r['course_id']]['end_time']}",
            courses[r["course_id"]]["name"], courses[r["course_id"]]["teacher_name"] or "", r["user_name"], r["phone"],
            STATUS.get(r["status"], r["status"]), r["charged"]] for r in res],
          [12, 13, 24, 12, 12, 14, 8, 6])
    sheet("課程統計", ["課程", "堂數", "總名額", "預約人次", "滿班率（%）", "出席"],
          [[c["name"], c["sessions"], c["capacity"], c["booked"], c["fill_rate"], c["attended"]] for c in data["courses"]],
          [28, 8, 10, 10, 12, 8])
    sheet("老師統計", ["老師", "授課堂數", "授課時數", "學員人次", "出席人次"],
          [[t["name"], t["sessions"], t["hours"], t["booked"], t["attended"]] for t in data["teachers"]],
          [16, 10, 10, 10, 10])
    sheet("會員排行", ["會員", "手機", "預約", "出席", "缺席"],
          [[m["name"], m["phone"], m["booked"], m["attended"], m["absent"]] for m in mem["top"]], [14, 14, 8, 8, 8])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def public(data: dict) -> dict:
    return {k: v for k, v in data.items() if not k.startswith("_")}


def parse_range(params, today: date) -> tuple[str, str]:
    try:
        end = date.fromisoformat(params.get("to") or today.isoformat())
        start = date.fromisoformat(params.get("from") or (end - timedelta(days=29)).isoformat())
    except ValueError:
        raise ValueError("日期格式不正確")
    if start > end or (end - start).days > 366:
        raise ValueError("期間最多一年，且開始日不能晚於結束日")
    return start.isoformat(), end.isoformat()

